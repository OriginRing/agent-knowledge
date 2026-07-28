import io
import re
import time
from datetime import datetime
from typing import Dict, List, Tuple


class DocumentService:
    """把结构化 Markdown 生成文件并上传到已经配置好的 OSS。"""

    MIME_TYPES = {
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "pdf": "application/pdf",
    }

    @staticmethod
    def _clean_inline_markdown(text: str) -> str:
        text = re.sub(r"`([^`]+)`", r"\1", text)
        text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
        text = re.sub(r"\*([^*]+)\*", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        return text.strip()

    @classmethod
    def build_docx(cls, markdown_text: str, title: str = "AI 分析报告") -> bytes:
        from docx import Document
        from docx.shared import Pt

        document = Document()
        normal_style = document.styles["Normal"]
        normal_style.font.name = "Microsoft YaHei"
        normal_style.font.size = Pt(10.5)
        has_heading = False
        for raw_line in markdown_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                document.add_heading(
                    cls._clean_inline_markdown(heading.group(2)),
                    level=min(len(heading.group(1)), 3),
                )
                has_heading = True
            elif re.match(r"^[-*+]\s+", line):
                document.add_paragraph(
                    cls._clean_inline_markdown(re.sub(r"^[-*+]\s+", "", line)),
                    style="List Bullet",
                )
            elif re.match(r"^\d+[.)]\s+", line):
                document.add_paragraph(
                    cls._clean_inline_markdown(re.sub(r"^\d+[.)]\s+", "", line)),
                    style="List Number",
                )
            elif "|" not in line or not line.startswith("|"):
                document.add_paragraph(cls._clean_inline_markdown(line))
        if not has_heading:
            if document.paragraphs:
                document.paragraphs[0].insert_paragraph_before(title, style="Title")
            else:
                document.add_heading(title, level=0)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _extract_markdown_tables(markdown_text: str) -> List[Tuple[str, List[List[str]]]]:
        tables: List[Tuple[str, List[List[str]]]] = []
        lines = markdown_text.splitlines()
        current_title = "数据"
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            heading = re.match(r"^#{1,6}\s+(.+)$", line)
            if heading:
                current_title = heading.group(1).strip()
            if (
                line.startswith("|")
                and index + 1 < len(lines)
                and re.match(r"^\|?[\s:|-]+\|?$", lines[index + 1].strip())
            ):
                rows = [[cell.strip() for cell in line.strip("|").split("|")]]
                index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append(
                        [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                    )
                    index += 1
                tables.append((current_title, rows))
                continue
            index += 1
        return tables

    @classmethod
    def build_xlsx(cls, markdown_text: str, title: str = "AI 生成表格") -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        workbook.remove(workbook.active)
        tables = cls._extract_markdown_tables(markdown_text)
        if not tables:
            tables = [(title, [["内容"], [cls._clean_inline_markdown(markdown_text)]])]
        used_names = set()
        for table_index, (table_title, rows) in enumerate(tables, start=1):
            base_name = re.sub(r"[:\\/?*\[\]]", "", table_title)[:31] or f"Sheet{table_index}"
            sheet_name = base_name
            suffix = 2
            while sheet_name in used_names:
                sheet_name = f"{base_name[:27]}-{suffix}"
                suffix += 1
            used_names.add(sheet_name)
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1677FF")
            for column in sheet.columns:
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
                sheet.column_dimensions[column[0].column_letter].width = width
            sheet.freeze_panes = "A2"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @classmethod
    def build_pptx(cls, markdown_text: str, title: str = "AI 生成演示文稿") -> bytes:
        from pptx import Presentation
        from pptx.util import Pt

        presentation = Presentation()
        title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        title_slide.shapes.title.text = title
        title_slide.placeholders[1].text = f"生成时间：{datetime.now():%Y-%m-%d %H:%M}"
        sections: List[Tuple[str, List[str]]] = []
        current_title = "内容概览"
        current_lines: List[str] = []
        for raw_line in markdown_text.splitlines():
            line = raw_line.strip()
            heading = re.match(r"^#{1,3}\s+(.+)$", line)
            if heading:
                if current_lines:
                    sections.append((current_title, current_lines))
                current_title = cls._clean_inline_markdown(heading.group(1))
                current_lines = []
            elif line:
                current_lines.append(cls._clean_inline_markdown(re.sub(r"^[-*+]\s+", "", line)))
        if current_lines:
            sections.append((current_title, current_lines))
        if not sections:
            sections = [(title, [cls._clean_inline_markdown(markdown_text)])]
        for section_title, bullets in sections:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = section_title
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, bullet in enumerate(bullets[:10]):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bullet
                paragraph.font.size = Pt(20)
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()

    @classmethod
    def build_pdf(cls, markdown_text: str, title: str = "AI 分析报告") -> bytes:
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        output = io.BytesIO()
        document = SimpleDocTemplate(output)
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "ChineseBody",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=10.5,
            leading=17,
            spaceAfter=8,
        )
        title_style = ParagraphStyle(
            "ChineseTitle",
            parent=body,
            fontSize=20,
            leading=26,
            alignment=TA_CENTER,
            spaceAfter=18,
        )
        heading_style = ParagraphStyle(
            "ChineseHeading",
            parent=body,
            fontSize=14,
            leading=20,
            spaceBefore=10,
            spaceAfter=8,
        )
        story = [Paragraph(cls._clean_inline_markdown(title), title_style)]
        for raw_line in markdown_text.splitlines():
            line = raw_line.strip()
            if not line:
                story.append(Spacer(1, 6))
                continue
            heading = re.match(r"^#{1,6}\s+(.+)$", line)
            text = cls._clean_inline_markdown(heading.group(1) if heading else line)
            text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(text, heading_style if heading else body))
        document.build(story)
        return output.getvalue()

    @classmethod
    def build_file(cls, content: str, artifact_format: str, title: str) -> bytes:
        builders = {
            "docx": cls.build_docx,
            "xlsx": cls.build_xlsx,
            "pptx": cls.build_pptx,
            "pdf": cls.build_pdf,
        }
        normalized = artifact_format.lower().lstrip(".")
        if normalized not in builders:
            raise ValueError(f"不支持的文件格式: {artifact_format}")
        return builders[normalized](content, title=title)

    @classmethod
    def create_and_upload(
        cls,
        content: str,
        title: str = "AI 生成文件",
        artifact_format: str = "docx",
    ) -> Dict:
        from services.file_service import FileService

        normalized = artifact_format.lower().lstrip(".")
        file_bytes = cls.build_file(content, normalized, title)
        filename = f"{title}-{datetime.now():%Y%m%d-%H%M%S}.{normalized}"
        result = FileService.upload_file(file_bytes, filename, object_prefix="generated")
        if result.get("code") != 0:
            raise RuntimeError(result.get("message", "文件上传失败"))
        data = result["data"]
        return {
            "id": f"artifact-{time.time_ns()}",
            "type": "file",
            "format": normalized,
            "mimeType": cls.MIME_TYPES[normalized],
            "name": data["filename"],
            "url": data["url"],
            "size": len(file_bytes),
        }
