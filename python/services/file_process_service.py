import base64
import hashlib
import io
import os
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import unquote, urlparse

import requests


class FileProcessService:
    IMAGE_EXTENSIONS = {
        ".jpg",
        ".jpeg",
        ".png",
        ".gif",
        ".bmp",
        ".webp",
        ".tif",
        ".tiff",
    }
    DOCUMENT_EXTENSIONS = {
        ".pdf",
        ".doc",
        ".docx",
        ".xls",
        ".xlsx",
        ".txt",
        ".md",
        ".ppt",
        ".pptx",
        ".html",
        ".htm",
        ".ofd",
    }
    SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS
    LEGACY_OFFICE_TARGETS = {
        ".doc": (".docx", "docx"),
        ".xls": (".xlsx", "xlsx"),
        ".ppt": (".pptx", "pptx"),
    }
    _pdfium_lock = threading.Lock()

    @classmethod
    def download_file_from_url(cls, url: str) -> bytes:
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"}:
            raise ValueError("仅支持 HTTP/HTTPS 文件地址")

        max_bytes = int(os.getenv("MAX_PARSE_FILE_BYTES", str(20 * 1024 * 1024)))
        with requests.get(
            url,
            timeout=float(os.getenv("FILE_DOWNLOAD_TIMEOUT_SECONDS", "30")),
            stream=True,
        ) as response:
            response.raise_for_status()
            content_length = int(response.headers.get("content-length") or 0)
            if content_length > max_bytes:
                raise ValueError(f"文件超过解析大小限制（{max_bytes // 1024 // 1024}MB）")
            chunks = []
            total = 0
            for chunk in response.iter_content(chunk_size=64 * 1024):
                if not chunk:
                    continue
                total += len(chunk)
                if total > max_bytes:
                    raise ValueError(
                        f"文件超过解析大小限制（{max_bytes // 1024 // 1024}MB）"
                    )
                chunks.append(chunk)
        return b"".join(chunks)

    @staticmethod
    def _section(
        text: str,
        source_kind: str,
        source_index: int,
        source_label: str,
        extraction_method: str,
    ) -> Dict[str, Any]:
        return {
            "text": (text or "").strip(),
            "sourceKind": source_kind,
            "sourceIndex": source_index,
            "sourceLabel": source_label,
            "extractionMethod": extraction_method,
        }

    @classmethod
    def _new_result(
        cls,
        sections: Optional[List[Dict[str, Any]]] = None,
        *,
        page_count: int = 0,
        image_count: int = 0,
        ocr_count: int = 0,
        warnings: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        usable_sections = [
            section for section in (sections or []) if section.get("text", "").strip()
        ]
        return {
            "sections": usable_sections,
            "pageCount": page_count,
            "imageCount": image_count,
            "ocrCount": ocr_count,
            "warnings": warnings or [],
            "content": cls._merge_sections(usable_sections),
        }

    @staticmethod
    def _merge_sections(sections: List[Dict[str, Any]]) -> str:
        if not sections:
            return ""
        if len(sections) == 1:
            return sections[0]["text"]
        return "\n\n".join(
            f"【{section['sourceLabel']}】\n{section['text']}" for section in sections
        )

    @classmethod
    def _ocr_image_bytes(
        cls, content: bytes, source_label: str
    ) -> Tuple[str, Optional[str]]:
        try:
            from services.ocr_service import OCRService

            text = OCRService.extract_text_from_image_bytes(content)
            if not text.strip():
                return "", f"{source_label} OCR 未识别到文本"
            return text.strip(), None
        except Exception as exc:
            return "", f"{source_label} OCR 失败: {exc}"

    @staticmethod
    def _dedupe_images(images: List[bytes]) -> List[bytes]:
        unique: List[bytes] = []
        seen = set()
        for content in images:
            digest = hashlib.sha256(content).hexdigest()
            if digest not in seen:
                seen.add(digest)
                unique.append(content)
        return unique

    @classmethod
    def _parse_docx(cls, content: bytes) -> Dict[str, Any]:
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = DocxDocument(io.BytesIO(content))
        sections: List[Dict[str, Any]] = []
        pending_blocks: List[str] = []
        seen_image_digests = set()
        warnings: List[str] = []
        image_count = 0
        ocr_count = 0
        text_section_index = 0

        def flush_text() -> None:
            nonlocal text_section_index
            if not pending_blocks:
                return
            text_section_index += 1
            label = (
                "文档正文"
                if text_section_index == 1
                else f"文档正文 {text_section_index}"
            )
            sections.append(
                cls._section(
                    "\n".join(pending_blocks),
                    "document",
                    text_section_index,
                    label,
                    "native",
                )
            )
            pending_blocks.clear()

        def append_images(element) -> None:
            nonlocal image_count, ocr_count
            relationship_ids = element.xpath(".//a:blip/@r:embed")
            for relationship_id in relationship_ids:
                part = document.part.related_parts.get(relationship_id)
                image_bytes = getattr(part, "blob", b"")
                if not image_bytes:
                    continue
                digest = hashlib.sha256(image_bytes).hexdigest()
                if digest in seen_image_digests:
                    continue
                seen_image_digests.add(digest)
                flush_text()
                image_count += 1
                label = f"文档图片 {image_count}"
                text, warning = cls._ocr_image_bytes(image_bytes, label)
                if text:
                    ocr_count += 1
                    sections.append(
                        cls._section(text, "image", image_count, label, "ocr")
                    )
                if warning:
                    warnings.append(warning)

        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if text:
                    pending_blocks.append(text)
                append_images(item._element)
            elif isinstance(item, Table):
                for row in item.rows:
                    row_text = "\t".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        pending_blocks.append(row_text)
                append_images(item._element)
        flush_text()

        # 兼容浮动对象等未出现在正文迭代结果中的图片关系。
        for relationship in document.part.rels.values():
            if (
                "image" not in relationship.reltype
                or getattr(relationship, "target_part", None) is None
            ):
                continue
            image_bytes = relationship.target_part.blob
            digest = hashlib.sha256(image_bytes).hexdigest()
            if digest in seen_image_digests:
                continue
            seen_image_digests.add(digest)
            image_count += 1
            label = f"文档图片 {image_count}"
            text, warning = cls._ocr_image_bytes(image_bytes, label)
            if text:
                ocr_count += 1
                sections.append(
                    cls._section(text, "image", image_count, label, "ocr")
                )
            if warning:
                warnings.append(warning)
        return cls._new_result(
            sections,
            image_count=image_count,
            ocr_count=ocr_count,
            warnings=warnings,
        )

    @classmethod
    def _parse_xlsx(cls, content: bytes) -> Dict[str, Any]:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sections: List[Dict[str, Any]] = []
        warnings: List[str] = []
        image_count = 0
        ocr_count = 0
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell) for cell in row if cell is not None]
                    if values:
                        rows.append("\t".join(values))
                if rows:
                    sections.append(
                        cls._section(
                            "\n".join(rows),
                            "sheet",
                            sheet_index,
                            f"工作表：{sheet.title}",
                            "native",
                        )
                    )

                for local_index, image in enumerate(
                    getattr(sheet, "_images", []), start=1
                ):
                    image_count += 1
                    label = f"工作表：{sheet.title} / 图片 {local_index}"
                    try:
                        image_bytes = image._data()
                    except Exception as exc:
                        warnings.append(f"{label} 读取失败: {exc}")
                        continue
                    text, warning = cls._ocr_image_bytes(image_bytes, label)
                    if text:
                        ocr_count += 1
                        sections.append(
                            cls._section(text, "sheet", sheet_index, label, "ocr")
                        )
                    if warning:
                        warnings.append(warning)
        finally:
            workbook.close()
        return cls._new_result(
            sections,
            image_count=image_count,
            ocr_count=ocr_count,
            warnings=warnings,
        )

    @classmethod
    def _parse_pptx(cls, content: bytes) -> Dict[str, Any]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(io.BytesIO(content))
        sections: List[Dict[str, Any]] = []
        warnings: List[str] = []
        image_count = 0
        ocr_count = 0
        for slide_index, slide in enumerate(presentation.slides, start=1):
            label = f"第 {slide_index} 张幻灯片"
            texts: List[str] = []
            slide_images: List[bytes] = []
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        row_text = "\t".join(
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        )
                        if row_text:
                            texts.append(row_text)
                elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                    texts.append(shape.text.strip())
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        slide_images.append(shape.image.blob)
                    except Exception as exc:
                        warnings.append(f"{label} 图片读取失败: {exc}")

            try:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text and notes_text not in texts:
                    texts.append(f"备注：{notes_text}")
            except (AttributeError, ValueError):
                pass

            if texts:
                sections.append(
                    cls._section(
                        "\n".join(texts), "slide", slide_index, label, "native"
                    )
                )
            for image_index, image_bytes in enumerate(
                cls._dedupe_images(slide_images), start=1
            ):
                image_count += 1
                image_label = f"{label} / 图片 {image_index}"
                text, warning = cls._ocr_image_bytes(image_bytes, image_label)
                if text:
                    ocr_count += 1
                    sections.append(
                        cls._section(
                            text, "slide", slide_index, image_label, "ocr"
                        )
                    )
                if warning:
                    warnings.append(warning)
        return cls._new_result(
            sections,
            page_count=len(presentation.slides),
            image_count=image_count,
            ocr_count=ocr_count,
            warnings=warnings,
        )

    @classmethod
    def _render_pdf_page(cls, content: bytes, page_index: int) -> bytes:
        import pypdfium2 as pdfium

        dpi = int(os.getenv("PDF_OCR_DPI", "180"))
        with cls._pdfium_lock:
            document = pdfium.PdfDocument(content)
            try:
                page = document[page_index]
                try:
                    bitmap = page.render(scale=dpi / 72)
                    try:
                        image = bitmap.to_pil()
                        output = io.BytesIO()
                        image.save(output, format="PNG")
                        return output.getvalue()
                    finally:
                        bitmap.close()
                finally:
                    page.close()
            finally:
                document.close()

    @classmethod
    def _parse_pdf(cls, content: bytes) -> Dict[str, Any]:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        threshold = int(os.getenv("PDF_OCR_MIN_TEXT_CHARS", "20"))
        sections: List[Dict[str, Any]] = []
        warnings: List[str] = []
        ocr_count = 0
        image_count = 0
        for page_index, page in enumerate(reader.pages, start=1):
            label = f"第 {page_index} 页"
            native_text = (page.extract_text() or "").strip()
            normalized_length = len("".join(native_text.split()))
            ocr_text = ""
            if normalized_length < threshold:
                image_count += 1
                try:
                    image_bytes = cls._render_pdf_page(content, page_index - 1)
                    ocr_text, warning = cls._ocr_image_bytes(image_bytes, label)
                    if warning:
                        warnings.append(warning)
                    if ocr_text:
                        ocr_count += 1
                except Exception as exc:
                    warnings.append(f"{label} 渲染或 OCR 失败: {exc}")

            text_parts = [part for part in (native_text, ocr_text) if part]
            if text_parts:
                sections.append(
                    cls._section(
                        "\n".join(text_parts),
                        "page",
                        page_index,
                        label,
                        "native+ocr" if native_text and ocr_text else (
                            "ocr" if ocr_text else "native"
                        ),
                    )
                )
        return cls._new_result(
            sections,
            page_count=len(reader.pages),
            image_count=image_count,
            ocr_count=ocr_count,
            warnings=warnings,
        )

    @classmethod
    def _parse_ofd(cls, content: bytes) -> Dict[str, Any]:
        try:
            from easyofd.ofd import OFD
        except ImportError as exc:
            raise RuntimeError("OFD 解析依赖 easyofd 未安装") from exc

        ofd = OFD()
        try:
            encoded = base64.b64encode(content).decode("ascii")
            ofd.read(encoded, save_xml=False)
            pdf_bytes = ofd.to_pdf()
            if not pdf_bytes:
                raise ValueError("OFD 转换未生成 PDF 内容")
        finally:
            try:
                ofd.del_data()
            except Exception:
                pass
        result = cls._parse_pdf(pdf_bytes)
        for section in result["sections"]:
            section["extractionMethod"] = f"ofd+{section['extractionMethod']}"
        result["content"] = cls._merge_sections(result["sections"])
        return result

    @classmethod
    def _parse_legacy_office(
        cls, content: bytes, filename: str, extension: str
    ) -> Dict[str, Any]:
        target_extension, convert_format = cls.LEGACY_OFFICE_TARGETS[extension]
        binary = os.getenv("LIBREOFFICE_BINARY") or shutil.which(
            "soffice"
        ) or shutil.which("libreoffice")
        if not binary:
            raise RuntimeError(
                f"解析 {extension} 需要安装 LibreOffice 或配置 LIBREOFFICE_BINARY"
            )

        with tempfile.TemporaryDirectory(prefix="agent-document-") as directory:
            root = Path(directory)
            input_path = root / Path(filename).name
            input_path.write_bytes(content)
            output_dir = root / "output"
            profile_dir = root / "profile"
            output_dir.mkdir()
            profile_dir.mkdir()
            command = [
                binary,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--convert-to",
                convert_format,
                "--outdir",
                str(output_dir),
                str(input_path),
            ]
            timeout = int(os.getenv("LIBREOFFICE_TIMEOUT_SECONDS", "60"))
            try:
                completed = subprocess.run(
                    command,
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=timeout,
                )
            except subprocess.TimeoutExpired as exc:
                raise RuntimeError(f"LibreOffice 转换超时（{timeout}秒）") from exc
            output_path = output_dir / f"{input_path.stem}{target_extension}"
            if completed.returncode != 0 or not output_path.exists():
                details = (completed.stderr or completed.stdout or "").strip()
                raise RuntimeError(
                    f"LibreOffice 转换失败{f'：{details}' if details else ''}"
                )
            result = cls.parse_document(
                output_path.read_bytes(), output_path.name
            )
        for section in result["sections"]:
            section["extractionMethod"] = (
                f"libreoffice+{section['extractionMethod']}"
            )
        result["content"] = cls._merge_sections(result["sections"])
        return result

    @classmethod
    def _parse_text(cls, content: bytes) -> str:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    @classmethod
    def parse_document(
        cls, content: bytes, filename: str, url: str = ""
    ) -> Dict[str, Any]:
        extension = os.path.splitext(filename)[1].lower()
        if extension not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"暂不支持解析该文件类型: {extension or '未知'}")

        if extension in cls.IMAGE_EXTENSIONS:
            text, warning = cls._ocr_image_bytes(content, "图片")
            sections = [cls._section(text, "image", 1, "图片", "ocr")] if text else []
            return cls._new_result(
                sections,
                image_count=1,
                ocr_count=1 if text else 0,
                warnings=[warning] if warning else [],
            )
        if extension == ".docx":
            return cls._parse_docx(content)
        if extension == ".xlsx":
            return cls._parse_xlsx(content)
        if extension == ".pptx":
            return cls._parse_pptx(content)
        if extension == ".pdf":
            return cls._parse_pdf(content)
        if extension == ".ofd":
            return cls._parse_ofd(content)
        if extension in cls.LEGACY_OFFICE_TARGETS:
            return cls._parse_legacy_office(content, filename, extension)
        if extension in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(cls._parse_text(content), "html.parser")
            for element in soup(["script", "style", "noscript"]):
                element.decompose()
            text = soup.get_text(separator="\n", strip=True)
        else:
            text = cls._parse_text(content)
        return cls._new_result(
            [cls._section(text, "document", 1, "文档正文", "native")]
        )

    @classmethod
    def extract_text(cls, content: bytes, filename: str, url: str = "") -> str:
        """兼容旧调用方：返回结构化解析结果的合并文本。"""
        return cls.parse_document(content, filename, url)["content"]

    @classmethod
    def extract_text_from_docx(cls, content: bytes) -> str:
        return cls._parse_docx(content)["content"]

    @classmethod
    def extract_text_from_xlsx(cls, content: bytes) -> str:
        return cls._parse_xlsx(content)["content"]

    @classmethod
    def extract_text_from_pptx(cls, content: bytes) -> str:
        return cls._parse_pptx(content)["content"]

    @classmethod
    def extract_text_from_pdf(cls, content: bytes) -> str:
        return cls._parse_pdf(content)["content"]

    @classmethod
    def extract_text_from_ofd(cls, content: bytes) -> str:
        return cls._parse_ofd(content)["content"]

    @classmethod
    def extract_text_from_html(cls, content: bytes) -> str:
        return cls.parse_document(content, "document.html")["content"]

    @classmethod
    def extract_text_from_txt(cls, content: bytes) -> str:
        return cls._parse_text(content)

    @classmethod
    def extract_text_from_markdown(cls, content: bytes) -> str:
        return cls._parse_text(content)

    @classmethod
    def extract_text_from_image(cls, url: str) -> str:
        from services.ocr_service import OCRService

        return OCRService.extract_text_from_image_url(url)

    @classmethod
    def process_files_for_ollama(cls, file_urls: List[str]) -> str:
        file_contents = [
            f"【文件内容】\n文件名：{item['filename']}\n文件内容：{item['content']}"
            for item in cls.process_files(file_urls)
            if item["status"] in {"success", "partial"} and item["content"].strip()
        ]
        return "\n\n".join(file_contents) + "\n\n" if file_contents else ""

    @classmethod
    def process_files_for_api(cls, file_urls: List[str]) -> List[Dict[str, Any]]:
        parsed_files = cls.process_files(file_urls)
        files = []
        for item in parsed_files:
            if item["isImage"]:
                files.append({"type": "image_url", "image_url": {"url": item["url"]}})
            if item["content"]:
                files.append(
                    {
                        "type": "text",
                        "text": f"【{item['filename']} 解析内容】\n{item['content']}",
                    }
                )
        return files

    @classmethod
    def process_files(cls, file_urls: List[str]) -> List[Dict[str, Any]]:
        results: List[Dict[str, Any]] = []
        for url in file_urls:
            path = unquote(urlparse(url).path)
            filename = os.path.basename(path) or "未命名文件"
            extension = os.path.splitext(filename)[1].lower()
            item: Dict[str, Any] = {
                "url": url,
                "filename": filename,
                "extension": extension,
                "isImage": extension in cls.IMAGE_EXTENSIONS,
                "content": "",
                "sections": [],
                "charCount": 0,
                "pageCount": 0,
                "imageCount": 0,
                "ocrCount": 0,
                "warnings": [],
                "status": "success",
                "error": None,
            }
            try:
                content = cls.download_file_from_url(url)
                parsed = cls.parse_document(content, filename, url)
                item.update(parsed)
                item["charCount"] = len(item["content"])
                if not item["content"].strip():
                    item["status"] = "error"
                    item["error"] = "未提取到可用内容"
                elif item["warnings"]:
                    item["status"] = "partial"
            except Exception as exc:
                item["status"] = "error"
                item["error"] = str(exc)
            results.append(item)
        return results
