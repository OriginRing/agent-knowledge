import json
import io
import os
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

import requests

from services.document_service import DocumentService
from services.file_process_service import FileProcessService
from services.file_service import FileService
from services.knowledge_service import KnowledgeService
from services.ocr_service import OCRService
from services.memory_service import format_memory_context
from services.skill_service import SkillService


class SkillServiceTest(unittest.TestCase):
    def tearDown(self):
        SkillService._cache = None

    def test_chart_visualization_skill_is_selected_for_chart_requests(self):
        for prompt in (
            "请生成销售额折线图",
            "把这些数据可视化展示出来",
            "绘制一张雷达图",
        ):
            with self.subTest(prompt=prompt):
                skill = SkillService.select_skill(prompt, [])
                self.assertIsNotNone(skill)
                self.assertEqual(skill.name, "chart-visualization")
                self.assertEqual(skill.kind, "hybrid")

    def test_chart_visualization_skill_can_be_selected_explicitly(self):
        skill = SkillService.select_skill(
            "展示这些数据",
            [],
            requested_skill="chart-visualization",
        )
        self.assertEqual(skill.name, "chart-visualization")
        self.assertIn("```vis line", skill.prompt)
        for chart_type in (
            "line",
            "area",
            "column",
            "bar",
            "pie",
            "scatter",
            "radar",
            "table",
        ):
            self.assertIn(f"`{chart_type}`", skill.prompt)

    def test_excel_table_request_does_not_select_chart_visualization(self):
        skill = SkillService.select_skill("请生成 Excel 表格", [])
        self.assertIsNone(skill)

    def test_builtin_image_skill_is_selected(self):
        skill = SkillService.select_skill(
            "请分析这张图并生成文档",
            ["https://example.com/demo.png?version=1"],
        )
        self.assertIsNotNone(skill)
        self.assertEqual(skill.name, "image-to-document")
        self.assertEqual(skill.artifact, "docx")

    def test_explicit_unknown_skill_raises(self):
        with self.assertRaisesRegex(ValueError, "技能不存在"):
            SkillService.select_skill("test", [], requested_skill="missing")

    def test_custom_skills_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            skill_dir = Path(directory) / "demo"
            skill_dir.mkdir()
            (skill_dir / "SKILL.md").write_text(
                "---\nname: demo\nintent_keywords: 测试\n---\n测试提示词",
                encoding="utf-8",
            )
            with patch.dict(os.environ, {"AGENT_SKILLS_DIR": directory}):
                skills = SkillService.load_skills(refresh=True)
            self.assertIn("demo", skills)
            self.assertEqual(skills["demo"].prompt, "测试提示词")

    def test_detects_multiple_artifact_formats_in_order(self):
        self.assertEqual(
            SkillService.detect_artifact_formats(
                "请生成 PDF、Word、Excel 和 PPT 四个文件"
            ),
            ["docx", "xlsx", "pptx", "pdf"],
        )
        self.assertEqual(
            SkillService.detect_artifact_formats("", "pdf,docx,pdf"),
            ["pdf", "docx"],
        )


class DocumentServiceTest(unittest.TestCase):
    def test_build_docx(self):
        content = DocumentService.build_docx(
            "# 图片分析报告\n\n- 识别内容\n\n1. 结论"
        )
        self.assertTrue(content.startswith(b"PK"))
        self.assertGreater(len(content), 1000)

    def test_create_and_upload_returns_artifact(self):
        with patch(
            "services.file_service.FileService.upload_file",
            return_value={
                "code": 0,
                "message": "success",
                "data": {
                    "url": "https://oss.example/report.docx",
                    "filename": "report.docx",
                },
            },
        ) as upload_mock:
            artifact = DocumentService.create_and_upload("# 报告\n内容")
        self.assertEqual(artifact["format"], "docx")
        self.assertEqual(artifact["url"], "https://oss.example/report.docx")
        self.assertEqual(upload_mock.call_args.kwargs["object_prefix"], "generated")

    def test_builds_all_supported_formats(self):
        markdown = "# 报告\n\n| 名称 | 数值 |\n| --- | --- |\n| A | 1 |\n\n## 结论\n- 正常"
        for artifact_format in ("docx", "xlsx", "pptx", "pdf"):
            with self.subTest(artifact_format=artifact_format):
                content = DocumentService.build_file(
                    markdown,
                    artifact_format,
                    "测试文件",
                )
                self.assertGreater(len(content), 500)
                if artifact_format != "pdf":
                    self.assertTrue(content.startswith(b"PK"))
                else:
                    self.assertTrue(content.startswith(b"%PDF"))

    @unittest.skipUnless(
        os.getenv("RUN_OSS_INTEGRATION") == "1",
        "设置 RUN_OSS_INTEGRATION=1 后执行真实 OSS 上传验证",
    )
    def test_uploads_all_formats_to_configured_oss(self):
        markdown = "# OSS 集成测试\n\n| 项目 | 状态 |\n| --- | --- |\n| 上传 | 正常 |"
        for artifact_format in ("docx", "xlsx", "pptx", "pdf"):
            with self.subTest(artifact_format=artifact_format):
                artifact = DocumentService.create_and_upload(
                    markdown,
                    title="OSS集成测试",
                    artifact_format=artifact_format,
                )
                self.assertTrue(artifact["url"].startswith("https://"))
                self.assertGreater(artifact["size"], 500)
                response = requests.head(artifact["url"], timeout=10)
                self.assertEqual(response.status_code, 200)


class FileServiceTest(unittest.TestCase):
    def test_upload_url_uses_date_directory_and_timestamped_filename(self):
        class FakeBucket:
            def __init__(self):
                self.object_key = ""

            def put_object(self, object_key, _content):
                self.object_key = object_key

        bucket = FakeBucket()
        with (
            patch.object(FileService, "get_bucket", return_value=bucket),
            patch.dict(
                os.environ,
                {
                    "OSS_BUCKET_NAME": "demo-bucket",
                    "OSS_ENDPOINT": "oss-cn-beijing.aliyuncs.com",
                },
            ),
        ):
            result = FileService.upload_file(
                b"content", "../季度 报告.docx"
            )

        self.assertEqual(result["code"], 0)
        self.assertEqual(result["data"]["filename"], "季度 报告.docx")
        key_parts = bucket.object_key.split("/")
        self.assertEqual(len(key_parts), 3)
        self.assertRegex(key_parts[1], r"^\d{8}$")
        self.assertRegex(
            key_parts[2], r"^季度 报告-\d+\.docx$"
        )
        self.assertRegex(
            result["data"]["url"],
            r"/\d{8}/%E5%AD%A3%E5%BA%A6%20%E6%8A%A5%E5%91%8A-\d+\.docx$",
        )


class MemoryContextTest(unittest.TestCase):
    def test_formats_nested_memos_response_without_duplicates(self):
        result = {
            "code": 0,
            "data": {
                "data": {
                    "memory_detail_list": [
                        {"memory_value": "用户喜欢简洁报告"},
                        {"memory_value": "用户喜欢简洁报告"},
                    ],
                    "preference_detail_list": [
                        {"content": "默认使用中文"},
                    ],
                }
            },
        }
        context = format_memory_context(result)
        self.assertEqual(context.count("用户喜欢简洁报告"), 1)
        self.assertIn("默认使用中文", context)

    def test_failed_memory_response_is_empty(self):
        self.assertEqual(format_memory_context({"code": -1}), "")


class FileProcessServiceTest(unittest.TestCase):
    def test_parsed_content_is_returned_for_node_output(self):
        parsed = FileProcessService._new_result(
            [
                FileProcessService._section(
                    "OCR 或文件解析正文", "image", 1, "图片", "ocr"
                )
            ],
            image_count=1,
            ocr_count=1,
        )
        with (
            patch.object(
                FileProcessService,
                "download_file_from_url",
                return_value=b"hello",
            ),
            patch.object(
                FileProcessService,
                "parse_document",
                return_value=parsed,
            ),
        ):
            result = FileProcessService.process_files(
                ["https://oss.example/path/demo.png?version=1"]
            )
        self.assertEqual(result[0]["filename"], "demo.png")
        self.assertEqual(result[0]["content"], "OCR 或文件解析正文")
        self.assertEqual(result[0]["charCount"], len("OCR 或文件解析正文"))
        self.assertTrue(result[0]["isImage"])
        self.assertEqual(result[0]["ocrCount"], 1)

    @staticmethod
    def _image_bytes(text: str = "OCR") -> bytes:
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (120, 50), "white")
        ImageDraw.Draw(image).text((8, 15), text, fill="black")
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()

    def test_docx_extracts_paragraph_table_and_embedded_image(self):
        from docx import Document
        from docx.shared import Inches

        document = Document()
        document.add_paragraph("正文内容")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "名称"
        table.cell(0, 1).text = "数值"
        document.add_picture(io.BytesIO(self._image_bytes()), width=Inches(1))
        output = io.BytesIO()
        document.save(output)

        with patch.object(
            FileProcessService,
            "_ocr_image_bytes",
            return_value=("图片文字", None),
        ):
            result = FileProcessService.parse_document(output.getvalue(), "demo.docx")

        self.assertIn("正文内容", result["content"])
        self.assertIn("名称\t数值", result["content"])
        self.assertIn("图片文字", result["content"])
        self.assertEqual(result["imageCount"], 1)
        self.assertEqual(result["ocrCount"], 1)

    def test_xlsx_extracts_sheets_and_embedded_images(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SpreadsheetImage
        from PIL import Image

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "数据"
        sheet.append(["项目", "结果"])
        sheet.append(["解析", "正常"])
        sheet.add_image(
            SpreadsheetImage(Image.open(io.BytesIO(self._image_bytes()))), "D1"
        )
        output = io.BytesIO()
        workbook.save(output)

        with patch.object(
            FileProcessService,
            "_ocr_image_bytes",
            return_value=("表格图片文字", None),
        ):
            result = FileProcessService.parse_document(output.getvalue(), "demo.xlsx")

        self.assertIn("项目\t结果", result["content"])
        self.assertIn("表格图片文字", result["content"])
        self.assertEqual(result["sections"][0]["sourceLabel"], "工作表：数据")
        self.assertEqual(result["imageCount"], 1)

    def test_pptx_extracts_text_table_and_embedded_image(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "幻灯片标题"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = "名称"
        table.cell(0, 1).text = "数值"
        slide.shapes.add_picture(io.BytesIO(self._image_bytes()), Inches(1), Inches(4))
        output = io.BytesIO()
        presentation.save(output)

        with patch.object(
            FileProcessService,
            "_ocr_image_bytes",
            return_value=("幻灯片图片文字", None),
        ):
            result = FileProcessService.parse_document(output.getvalue(), "demo.pptx")

        self.assertIn("幻灯片标题", result["content"])
        self.assertIn("名称\t数值", result["content"])
        self.assertIn("幻灯片图片文字", result["content"])
        self.assertEqual(result["pageCount"], 1)

    def test_pdf_only_ocrs_pages_without_native_text(self):
        from reportlab.pdfgen import canvas

        output = io.BytesIO()
        pdf = canvas.Canvas(output)
        pdf.drawString(72, 760, "This page has enough native text for extraction.")
        pdf.showPage()
        pdf.showPage()
        pdf.save()

        with (
            patch.object(
                FileProcessService,
                "_render_pdf_page",
                return_value=self._image_bytes(),
            ) as render_mock,
            patch.object(
                FileProcessService,
                "_ocr_image_bytes",
                return_value=("扫描页文字", None),
            ),
        ):
            result = FileProcessService.parse_document(output.getvalue(), "demo.pdf")

        self.assertEqual(render_mock.call_count, 1)
        self.assertEqual(result["ocrCount"], 1)
        self.assertIn("扫描页文字", result["content"])

    def test_legacy_office_uses_isolated_libreoffice_conversion(self):
        from docx import Document

        def fake_run(command, **_):
            output_dir = Path(command[command.index("--outdir") + 1])
            input_path = Path(command[-1])
            document = Document()
            document.add_paragraph("旧版 Word 正文")
            document.save(output_dir / f"{input_path.stem}.docx")
            return subprocess.CompletedProcess(command, 0, "", "")

        with (
            patch("services.file_process_service.shutil.which", return_value="/fake/soffice"),
            patch("services.file_process_service.subprocess.run", side_effect=fake_run),
        ):
            result = FileProcessService.parse_document(b"legacy", "demo.doc")

        self.assertIn("旧版 Word 正文", result["content"])
        self.assertTrue(
            result["sections"][0]["extractionMethod"].startswith("libreoffice+")
        )

    def test_ofd_conversion_reuses_pdf_page_pipeline(self):
        class FakeOFD:
            def read(self, encoded, save_xml=False):
                self.content = encoded
                self.save_xml = save_xml

            def to_pdf(self):
                return b"%PDF-fake"

            def del_data(self):
                return None

        easyofd_module = types.ModuleType("easyofd")
        easyofd_module.__path__ = []
        ofd_module = types.ModuleType("easyofd.ofd")
        ofd_module.OFD = FakeOFD
        parsed_pdf = FileProcessService._new_result(
            [
                FileProcessService._section(
                    "OFD 页面文字", "page", 1, "第 1 页", "ocr"
                )
            ],
            page_count=1,
            ocr_count=1,
        )
        with (
            patch.dict(
                sys.modules,
                {"easyofd": easyofd_module, "easyofd.ofd": ofd_module},
            ),
            patch.object(
                FileProcessService, "_parse_pdf", return_value=parsed_pdf
            ) as parse_pdf,
        ):
            result = FileProcessService.parse_document(b"ofd", "demo.ofd")

        parse_pdf.assert_called_once_with(b"%PDF-fake")
        self.assertEqual(result["sections"][0]["extractionMethod"], "ofd+ocr")

    def test_ocr_warning_marks_file_as_partial(self):
        parsed = FileProcessService._new_result(
            [
                FileProcessService._section(
                    "已有正文", "document", 1, "文档正文", "native"
                )
            ],
            image_count=1,
            warnings=["图片 1 OCR 失败"],
        )
        with (
            patch.object(
                FileProcessService, "download_file_from_url", return_value=b"content"
            ),
            patch.object(FileProcessService, "parse_document", return_value=parsed),
        ):
            result = FileProcessService.process_files(
                ["https://oss.example/demo.docx"]
            )[0]
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["content"], "已有正文")


class OCRServiceTest(unittest.TestCase):
    def tearDown(self):
        OCRService._client = None

    def test_image_bytes_are_normalized_to_data_url(self):
        image_bytes = FileProcessServiceTest._image_bytes()
        with patch.object(
            OCRService, "_recognize", return_value="识别结果"
        ) as recognize:
            text = OCRService.extract_text_from_image_bytes(image_bytes)
        self.assertEqual(text, "识别结果")
        self.assertTrue(recognize.call_args.args[0].startswith("data:image/jpeg;base64,"))

    def test_client_configuration_comes_from_environment(self):
        OCRService._client = None
        env = {
            "QWEN_API_KEY": "test-key",
            "QWEN_OCR_BASE_URL": "https://ocr.example.com/v1",
            "OCR_TIMEOUT_SECONDS": "12.5",
        }
        with patch.dict(os.environ, env, clear=True), patch(
            "services.ocr_service.OpenAI"
        ) as openai:
            OCRService.get_client()
        openai.assert_called_once_with(
            api_key="test-key",
            base_url="https://ocr.example.com/v1",
            timeout=12.5,
        )

    def test_ocr_model_comes_from_environment(self):
        completion = types.SimpleNamespace(
            choices=[
                types.SimpleNamespace(
                    message=types.SimpleNamespace(content="识别结果")
                )
            ]
        )
        env = {
            "QWEN_OCR_MODEL": "custom-ocr-model",
            "OCR_MAX_RETRIES": "0",
        }
        with patch.dict(os.environ, env, clear=True), patch.object(
            OCRService, "get_client"
        ) as get_client:
            get_client.return_value.chat.completions.create.return_value = completion
            result = OCRService._recognize("data:image/jpeg;base64,demo")

        self.assertEqual(result, "识别结果")
        self.assertEqual(
            get_client.return_value.chat.completions.create.call_args.kwargs["model"],
            "custom-ocr-model",
        )


class KnowledgeSectionTest(unittest.TestCase):
    def tearDown(self):
        KnowledgeService._embeddings = None

    def test_embedding_model_comes_from_environment(self):
        KnowledgeService._embeddings = None
        with patch.dict(
            os.environ,
            {"KNOWLEDGE_EMBEDDING_MODEL": "custom-embedding-model"},
            clear=True,
        ), patch(
            "services.knowledge_service.OllamaEmbeddings"
        ) as embeddings:
            KnowledgeService.get_embeddings()

        embeddings.assert_called_once_with(model="custom-embedding-model")

    def test_sections_are_chunked_without_crossing_source_boundaries(self):
        documents = KnowledgeService.split_sections(
            [
                {
                    "text": "第一页内容",
                    "sourceKind": "page",
                    "sourceIndex": 1,
                    "sourceLabel": "第 1 页",
                    "extractionMethod": "native",
                },
                {
                    "text": "第二页扫描内容",
                    "sourceKind": "page",
                    "sourceIndex": 2,
                    "sourceLabel": "第 2 页",
                    "extractionMethod": "ocr",
                },
            ],
            "demo.pdf",
            "https://oss.example/demo.pdf",
            "demo-1",
            1,
        )
        self.assertEqual(len(documents), 2)
        self.assertEqual(documents[0].metadata["source_index"], 1)
        self.assertEqual(documents[1].metadata["source_label"], "第 2 页")
        self.assertEqual(documents[1].metadata["extraction_method"], "ocr")


class ChatPipelineTest(unittest.IsolatedAsyncioTestCase):
    @staticmethod
    def _config():
        return {
            "agent_code": "000001",
            "agent_name": "测试智能体",
            "model_type": "ollama",
            "model_name": "test",
            "base_url": "",
            "support_file": True,
            "support_think": True,
            "support_connect": True,
            "support_knowledge": True,
        }

    async def test_explicit_chart_skill_injects_gpt_vis_system_prompt(self):
        from agent.agent_service import AgentService

        config = self._config()
        captured_messages = []

        async def fake_model_stream(
            _config, _text, _thinking, _connect, base_messages, context
        ):
            captured_messages.extend(base_messages)
            yield AgentService._event(
                config,
                event="message",
                content=(
                    "月度销售额\n\n"
                    "```vis line\n"
                    "data\n"
                    "  - time 2024-01\n"
                    "    value 120\n"
                    "```\n\n"
                    "销售额为 120。"
                ),
                **context,
            )

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(
                AgentService, "_chat_stream_ollama", side_effect=fake_model_stream
            ),
            patch("services.memory_service.is_memory_enabled", return_value=False),
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="展示这些数据",
                    skill="chart-visualization",
                )
            ]

        self.assertTrue(
            any(
                "GPT-Vis 图表生成" in item["content"]
                and "```vis line" in item["content"]
                for item in captured_messages
                if item["role"] == "system"
            )
        )
        skill_node = next(
            item["node"]
            for item in chunks
            if item.get("node", {}).get("name") == "skill"
        )
        self.assertIn("chart-visualization", skill_node["details"]["skills"])
        self.assertEqual(chunks[-1]["skill"], "chart-visualization")

    async def test_multiple_generated_files_use_one_file_node(self):
        from agent.agent_service import AgentService

        config = self._config()

        async def fake_model_stream(
            _config, _text, _thinking, _connect, _base_messages, context
        ):
            yield AgentService._event(
                config,
                event="message",
                content="# 报告\n内容",
                think_message="正在组织报告",
                **context,
            )
            yield AgentService._step(
                config,
                "model_call",
                "completed",
                "模型调用完成",
                node_id="model-call-1",
                node_kind="model",
                **context,
            )

        def execute_skill(name, **kwargs):
            self.assertEqual(name, "artifact-generator")
            artifact_format = kwargs["artifact_format"]
            return {
                "id": f"file-{artifact_format}",
                "type": "file",
                "format": artifact_format,
                "mimeType": "application/octet-stream",
                "name": f"报告.{artifact_format}",
                "url": f"https://oss.example/report.{artifact_format}",
                "size": 1024,
            }

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(
                AgentService, "_chat_stream_ollama", side_effect=fake_model_stream
            ),
            patch(
                "services.skill_service.SkillService.execute",
                side_effect=execute_skill,
            ),
            patch("services.memory_service.is_memory_enabled", return_value=True),
            patch(
                "services.memory_service.search_memory",
                return_value={"code": 0, "data": {}},
            ),
            patch(
                "services.memory_service.add_memory",
                return_value={"code": 0, "message": "success"},
            ) as add_memory_mock,
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="生成 Word 和 PDF",
                    thinking=True,
                    username="10001",
                )
            ]

        skill_node = next(
            item["node"]
            for item in chunks
            if item.get("node", {}).get("name") == "skill"
        )
        self.assertIn(
            "artifact-generator", skill_node["details"]["skills"]
        )
        file_nodes = [
            item["node"]
            for item in chunks
            if item.get("node", {}).get("id") == "skill-artifact-generator"
        ]
        self.assertEqual(file_nodes[-1]["kind"], "file")
        self.assertEqual(file_nodes[-1]["status"], "success")
        self.assertEqual(
            file_nodes[-1]["fileUrl"],
            "https://oss.example/report.docx,https://oss.example/report.pdf",
        )
        self.assertEqual(len(file_nodes[-1]["details"]["files"]), 2)
        artifact_event = next(
            item for item in chunks if item.get("event") == "artifact"
        )
        self.assertEqual(artifact_event["fileUrl"], file_nodes[-1]["fileUrl"])
        self.assertEqual(chunks[-1]["fileUrl"], file_nodes[-1]["fileUrl"])
        reasoning_nodes = [
            item["node"]
            for item in chunks
            if item.get("node", {}).get("id") == "model-call-1"
            and item["node"].get("details", {}).get("reasoning")
        ]
        self.assertEqual(
            reasoning_nodes[-1]["details"]["reasoning"], "正在组织报告"
        )
        message_content = "".join(
            item.get("content", "")
            for item in chunks
            if item.get("event") == "message"
        )
        self.assertEqual(message_content, "# 报告\n内容")
        self.assertNotIn("https://oss.example", message_content)
        add_memory_mock.assert_not_called()

    async def test_partial_file_failure_keeps_successful_url(self):
        from agent.agent_service import AgentService

        config = self._config()

        async def fake_model_stream(
            _config, _text, _thinking, _connect, _base_messages, context
        ):
            yield AgentService._event(
                config, event="message", content="报告正文", **context
            )

        def execute_skill(_name, **kwargs):
            if kwargs["artifact_format"] == "pdf":
                raise RuntimeError("PDF 上传失败")
            return {
                "id": "file-docx",
                "type": "file",
                "format": "docx",
                "mimeType": "application/octet-stream",
                "name": "报告.docx",
                "url": "https://oss.example/report.docx",
                "size": 1024,
            }

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(
                AgentService, "_chat_stream_ollama", side_effect=fake_model_stream
            ),
            patch(
                "services.skill_service.SkillService.execute",
                side_effect=execute_skill,
            ),
            patch("services.memory_service.is_memory_enabled", return_value=False),
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="生成 Word 和 PDF",
                )
            ]

        file_node = [
            item["node"]
            for item in chunks
            if item.get("node", {}).get("id") == "skill-artifact-generator"
        ][-1]
        self.assertEqual(file_node["status"], "error")
        self.assertEqual(file_node["fileUrl"], "https://oss.example/report.docx")
        self.assertEqual(len(file_node["details"]["errors"]), 1)

    async def test_all_file_failures_return_empty_url(self):
        from agent.agent_service import AgentService

        config = self._config()

        async def fake_model_stream(
            _config, _text, _thinking, _connect, _base_messages, context
        ):
            yield AgentService._event(
                config, event="message", content="报告正文", **context
            )

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(
                AgentService, "_chat_stream_ollama", side_effect=fake_model_stream
            ),
            patch(
                "services.skill_service.SkillService.execute",
                side_effect=RuntimeError("OSS 不可用"),
            ),
            patch("services.memory_service.is_memory_enabled", return_value=False),
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="生成 Word 和 PDF",
                )
            ]

        file_node = [
            item["node"]
            for item in chunks
            if item.get("node", {}).get("id") == "skill-artifact-generator"
        ][-1]
        self.assertEqual(file_node["status"], "error")
        self.assertEqual(file_node["fileUrl"], "")
        self.assertEqual(file_node["details"]["files"], [])
        self.assertEqual(len(file_node["details"]["errors"]), 2)
        self.assertFalse(any(item.get("event") == "artifact" for item in chunks))

    async def test_memory_is_injected_without_writing_conversation(self):
        from agent.agent_service import AgentService

        config = {
            "agent_code": "000001",
            "agent_name": "测试智能体",
            "model_type": "ollama",
            "model_name": "test",
            "base_url": "",
            "support_file": False,
            "support_think": False,
            "support_connect": False,
            "support_knowledge": False,
        }
        captured_messages = []

        async def fake_model_stream(
            _config, _text, _thinking, _connect, base_messages, context
        ):
            captured_messages.extend(base_messages)
            yield AgentService._event(
                config, event="message", content="回答", **context
            )
            yield AgentService._step(
                config, "model_call", "completed", "模型调用完成", **context
            )

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(
                AgentService,
                "_chat_stream_ollama",
                side_effect=fake_model_stream,
            ),
            patch(
                "services.memory_service.is_memory_enabled",
                return_value=True,
            ),
            patch(
                "services.memory_service.search_memory",
                return_value={
                    "code": 0,
                    "data": {"memory_detail_list": [{"memory_value": "偏好中文"}]},
                },
            ),
            patch(
                "services.memory_service.add_memory",
                return_value={"code": 0, "message": "success"},
            ) as add_memory_mock,
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="你好",
                    username="10001",
                    session_id=None,
                )
            ]

        self.assertTrue(any("偏好中文" in item["content"] for item in captured_messages))
        self.assertTrue(any(item.get("event") == "node" for item in chunks))
        self.assertEqual(chunks[-1]["event"], "done")
        self.assertTrue(chunks[-1]["memory"])
        self.assertFalse(
            any(
                item.get("node", {}).get("name") == "memory_write"
                for item in chunks
            )
        )
        add_memory_mock.assert_not_called()

    async def test_forced_search_skills_emit_nodes_and_feed_model_context(self):
        from agent.agent_service import AgentService

        config = {
            "agent_code": "000001",
            "agent_name": "测试智能体",
            "model_type": "ollama",
            "model_name": "test",
            "base_url": "",
            "support_file": True,
            "support_think": False,
            "support_connect": True,
            "support_knowledge": True,
        }
        execution_order = []
        captured_messages = []

        def execute_skill(name, **_):
            execution_order.append(name)
            if name == "web-search":
                return {"context": "联网结果", "items": [{"title": "网页"}]}
            return {
                "context": "知识内容",
                "items": [{"fileId": "knowledge-1"}],
            }

        async def fake_model_stream(
            _config, _text, _thinking, _connect, base_messages, context
        ):
            captured_messages.extend(base_messages)
            yield AgentService._event(
                config, event="message", content="综合回答", **context
            )
            yield AgentService._step(
                config,
                "model_call",
                "completed",
                "模型调用完成",
                node_id="model-call-1",
                node_kind="model",
                **context,
            )

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(AgentService, "_chat_stream_ollama", side_effect=fake_model_stream),
            patch("services.skill_service.SkillService.execute", side_effect=execute_skill),
            patch("services.memory_service.is_memory_enabled", return_value=False),
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="查询",
                    skills=["web-search", "knowledge-search"],
                )
            ]

        self.assertEqual(execution_order, ["web-search", "knowledge-search"])
        self.assertTrue(any("联网结果" in item["content"] for item in captured_messages))
        self.assertTrue(any("知识内容" in item["content"] for item in captured_messages))
        node_ids = {item["node"]["id"] for item in chunks if item.get("node")}
        self.assertIn("skill-web-search", node_ids)
        self.assertIn("skill-knowledge-search", node_ids)
        self.assertIn("model-call-1", node_ids)

    async def test_skill_failure_does_not_prevent_model_answer(self):
        from agent.agent_service import AgentService

        config = {
            "agent_code": "000001",
            "agent_name": "测试智能体",
            "model_type": "ollama",
            "model_name": "test",
            "base_url": "",
            "support_file": False,
            "support_think": False,
            "support_connect": True,
            "support_knowledge": False,
        }

        async def fake_model_stream(
            _config, _text, _thinking, _connect, _base_messages, context
        ):
            yield AgentService._event(
                config, event="message", content="降级回答", **context
            )
            yield AgentService._step(
                config,
                "model_call",
                "completed",
                "模型调用完成",
                node_id="model-call-1",
                node_kind="model",
                **context,
            )

        with (
            patch.object(AgentService, "get_agent_config", return_value=config),
            patch.object(AgentService, "_chat_stream_ollama", side_effect=fake_model_stream),
            patch("services.skill_service.SkillService.execute", side_effect=RuntimeError("搜索不可用")),
            patch("services.memory_service.is_memory_enabled", return_value=False),
        ):
            chunks = [
                json.loads(chunk)
                async for chunk in AgentService.chat_stream(
                    agent_code="000001",
                    text="查询",
                    skills=["web-search"],
                )
            ]
        self.assertTrue(any(item.get("content") == "降级回答" for item in chunks))
        self.assertTrue(
            any(
                item.get("node", {}).get("id") == "skill-web-search"
                and item["node"]["status"] == "error"
                for item in chunks
            )
        )
        self.assertEqual(chunks[-1]["event"], "done")


if __name__ == "__main__":
    unittest.main()
